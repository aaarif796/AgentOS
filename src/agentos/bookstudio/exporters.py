from __future__ import annotations

import re
from enum import StrEnum
from pathlib import Path

from agentos.bookstudio.schema import BookSpec

# ---------------------------------------------------------------------------
# Markdown assembler (canonical intermediate format for every exporter)
# ---------------------------------------------------------------------------


def render_markdown(book: BookSpec, toc: bool = True, references: bool = True) -> str:
    lines: list[str] = []
    lines.append(f"# {book.title}")
    if book.subtitle:
        lines.append(f"## {book.subtitle}")
    if book.audience:
        lines.append(f"*Audience: {book.audience}*")
    lines.append("")
    if book.front_matter:
        lines.append(book.front_matter.strip())
        lines.append("")
    if toc:
        lines.append("## Table of Contents")
        for i, ch in enumerate(book.chapters, start=1):
            lines.append(f"{i}. {ch.title}")
        lines.append("")
    for i, ch in enumerate(book.chapters, start=1):
        lines.append("---")
        lines.append(f"## Chapter {i}: {ch.title}")
        body = ch.draft.strip() or "*pending*"
        lines.append(body)
        lines.append("")
    if book.back_matter:
        lines.append("---")
        lines.append(book.back_matter.strip())
        lines.append("")
    if references:
        source_list = _collect_sources(book)
        if source_list:
            lines.append("## References")
            seen: set[str] = set()
            for src in source_list:
                if src and src not in seen:
                    lines.append(f"- {src}")
                    seen.add(src)
            lines.append("")
    return "\n".join(lines)


def _collect_sources(book: BookSpec) -> list[str]:
    out: list[str] = []
    for ch in book.chapters:
        out.extend(ch.sources)
    return out


# ---------------------------------------------------------------------------
# DOCX exporter (python-docx)
# ---------------------------------------------------------------------------


def export_docx(markdown: str, out_path: Path) -> Path:
    from docx import Document
    from docx.shared import Inches, Pt

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    for line in markdown.splitlines():
        raw = _inline(line)
        stripped = line.strip()
        if stripped.startswith("# "):
            doc.add_heading(raw.lstrip("# ").strip(), level=0)
        elif stripped.startswith("## "):
            doc.add_heading(raw.lstrip("# ").strip(), level=1)
        elif stripped.startswith("### "):
            doc.add_heading(raw.lstrip("# ").strip(), level=2)
        elif stripped == "---" or stripped == "***":
            doc.add_paragraph()
        elif stripped.startswith(("1. ", "2. ", "- ", "* ")):
            text = _lstrip_marker(raw)
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(text)
        elif stripped.startswith("```"):
            continue  # code fence markers dropped; content lines flow below
        elif stripped:
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(raw)
            font = run.font
            font.size = Pt(11)
            font.name = "Calibri"
            if stripped.startswith("*") and stripped.endswith("*"):
                font.italic = True
            if stripped.startswith("**") and stripped.endswith("**"):
                font.bold = True

    doc.save(str(out_path))
    return out_path


# ---------------------------------------------------------------------------
# PDF exporter (reportlab) with Unicode/TTF font auto-detection
# ---------------------------------------------------------------------------

_FONT_CANDIDATES = [
    "DejaVuSans.ttf",
    "DejaVuSansMono.ttf",
    "Arial.ttf",
    "Calibri.ttf",
    "Vera.ttf",
    "NotoSans-Regular.ttf",
]


def _find_base_font() -> tuple[str, Path] | None:
    """Find a Unicode-capable TrueType font so non-Latin text renders in PDFs."""
    search_paths = [
        Path("C:/Windows/Fonts"),
        Path("/usr/share/fonts/truetype/dejavu"),
        Path("/usr/share/fonts"),
    ]
    try:
        import reportlab

        search_paths.append(Path(reportlab.__file__).parent / "fonts")
    except Exception:
        pass
    for base in search_paths:
        for name in _FONT_CANDIDATES:
            candidate = base / name
            if candidate.exists():
                return name.rsplit(".", 1)[0], candidate
    return None


def export_pdf(markdown: str, out_path: Path) -> Path:
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Paragraph,
        Preformatted,
        SimpleDocTemplate,
        Spacer,
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    font_name = "Helvetica"
    found = _find_base_font()
    if found:
        family, font_path = found
        try:
            pdfmetrics.registerFont(TTFont(family, str(font_path)))
            font_name = family
        except Exception:
            font_name = "Helvetica"

    base = getSampleStyleSheet()
    styles = {
        "h1": ParagraphStyle(
            "h1", parent=base["Heading1"], fontName=font_name, fontSize=26, leading=30
        ),
        "h2": ParagraphStyle(
            "h2", parent=base["Heading2"], fontName=font_name, fontSize=18, leading=24
        ),
        "h3": ParagraphStyle(
            "h3", parent=base["Heading3"], fontName=font_name, fontSize=14, leading=18
        ),
        "body": ParagraphStyle(
            "body", parent=base["BodyText"], fontName=font_name, alignment=TA_LEFT
        ),
        "bullet": ParagraphStyle(
            "bullet",
            parent=base["BodyText"],
            fontName=font_name,
            leftIndent=0.3 * inch,
            bulletIndent=0.1 * inch,
        ),
        "code": ParagraphStyle(
            "code", fontName="Courier", fontSize=9, leading=12, backColor="#f4f4f4"
        ),
    }

    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=letter,
        leftMargin=1.25 * inch,
        rightMargin=1.25 * inch,
        topMargin=1 * inch,
        bottomMargin=1 * inch,
        title=out_path.stem,
    )
    flow: list[object] = []
    code_buffer: list[str] = []
    for line in markdown.splitlines():
        stripped = line.strip()
        if stripped.startswith("```"):
            if code_buffer:
                flow.append(Preformatted("\n".join(code_buffer), styles["code"]))
                flow.append(Spacer(1, 6))
                code_buffer = []
            continue
        if code_buffer:
            code_buffer.append(line)
            continue
        if stripped.startswith("# "):
            flow.append(Paragraph(_esc(stripped[2:]), styles["h1"]))
        elif stripped.startswith("## "):
            flow.append(Paragraph(_esc(stripped[3:]), styles["h2"]))
        elif stripped.startswith("### "):
            flow.append(Paragraph(_esc(stripped[4:]), styles["h3"]))
        elif stripped == "---" or stripped == "***":
            flow.append(Spacer(1, 8))
        elif stripped.startswith(("- ", "* ")):
            flow.append(Paragraph("• " + _esc(_lstrip_marker(stripped)), styles["bullet"]))
        elif re.match(r"^\d+\.\s", stripped):
            flow.append(Paragraph(_esc(stripped), styles["bullet"]))
        elif stripped:
            flow.append(Paragraph(_esc(stripped), styles["body"]))
            flow.append(Spacer(1, 3))
    if code_buffer:
        flow.append(Preformatted("\n".join(code_buffer), styles["code"]))

    doc.build(flow)
    return out_path


# ---------------------------------------------------------------------------
# ODF exporter (odfpy)
# ---------------------------------------------------------------------------


def export_odf(markdown: str, out_path: Path) -> Path:
    from odf.opendocument import OpenDocumentText
    from odf.style import ParagraphProperties, Style
    from odf.text import H, P

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc_text = OpenDocumentText()

    body_style = Style(name="Body", family="paragraph")
    body_style.addElement(ParagraphProperties(linenumber="true"))
    doc_text.automaticstyles.addElement(body_style)

    for level in range(1, 4):
        style = Style(name=f"H{level}", family="paragraph", parentstylename=f"Heading {level}")
        doc_text.styles.addElement(style)

    text = doc_text.text
    for line in markdown.splitlines():
        stripped = line.strip()
        m = _HEADING_RE.match(stripped)
        if m:
            level = len(m.group(1))
            heading = H(outlinelevel=level, text=_inline(m.group(2)))
            text.addElement(heading)
        elif stripped.startswith(("1. ", "2. ", "- ", "* ")):
            p = P(stylename="Body", text=_inline(_lstrip_marker(stripped)))
            text.addElement(p)
        elif stripped == "---" or stripped == "***":
            continue
        elif stripped:
            p = P(stylename="Body", text=_inline(stripped))
            text.addElement(p)

    doc_text.save(str(out_path))
    return out_path


class ExportFormat(StrEnum):
    MD = "md"
    DOCX = "docx"
    PDF = "pdf"
    ODF = "odf"


def export_book(
    book: BookSpec, out_dir: Path, formats: list[ExportFormat] | None = None
) -> dict[str, str]:
    """Compile a book into the requested formats. Returns {format: absolute path}."""
    out_dir = out_dir.resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    markdown = render_markdown(book)
    requested = formats or [ExportFormat.MD, ExportFormat.DOCX, ExportFormat.PDF, ExportFormat.ODF]
    artifacts: dict[str, str] = {}
    slug = re.sub(r"[^a-z0-9]+", "-", book.title.lower()).strip("-")[:60] or "book"
    for fmt in requested:
        target = out_dir / f"{slug}.{fmt.value}"
        if fmt == ExportFormat.MD:
            target.write_text(markdown, encoding="utf-8")
        elif fmt == ExportFormat.DOCX:
            export_docx(markdown, target)
        elif fmt == ExportFormat.PDF:
            export_pdf(markdown, target)
        elif fmt == ExportFormat.ODF:
            export_odf(markdown, target)
        artifacts[fmt.value] = str(target)
    return artifacts


def _esc(text: str) -> str:
    """Escape XML special chars that reportlab's Paragraph parser needs."""
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ---------------------------------------------------------------------------
# Shared markdown helpers
# ---------------------------------------------------------------------------


def _inline(text: str) -> str:
    return text.replace("**", "").replace("*", "").replace("`", "").replace("_", "").strip()


def _lstrip_marker(text: str) -> str:
    return re.sub(r"^\s*(?:\d+\.|-|\*)\s+", "", text, count=1)


_HEADING_RE = re.compile(r"^(#{1,3})\s+(.*)$")
